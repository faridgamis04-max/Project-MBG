"""Convert generated Markdown (+ optional Kop Surat meta) into Word / PDF / Excel bytes."""
import io
import re
from typing import List, Dict, Optional

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Font as XLFont, Alignment as XLAlign, Border, Side, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable,
)


# ---------------------------------------------------------------------------
# Markdown parsing -> list of blocks
# ---------------------------------------------------------------------------
def _clean_inline(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text.strip()


def _is_table_sep(line: str) -> bool:
    s = line.strip().strip("|")
    cells = [c.strip() for c in s.split("|")]
    return len(cells) > 0 and all(re.fullmatch(r":?-{2,}:?", c or "-") or set(c) <= set("-: ") and "-" in c for c in cells if c != "")


def parse_markdown(md: str) -> List[Dict]:
    lines = md.replace("\r\n", "\n").split("\n")
    blocks: List[Dict] = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        # Table detection: a line with | followed by a separator line
        if stripped.startswith("|") or ("|" in stripped and i + 1 < n and "|" in lines[i + 1] and _is_table_sep(lines[i + 1])):
            if i + 1 < n and _is_table_sep(lines[i + 1]):
                headers = [_clean_inline(c) for c in stripped.strip("|").split("|")]
                rows = []
                j = i + 2
                while j < n and "|" in lines[j] and lines[j].strip():
                    cells = [_clean_inline(c) for c in lines[j].strip().strip("|").split("|")]
                    # normalize cell count
                    while len(cells) < len(headers):
                        cells.append("")
                    rows.append(cells[: len(headers)])
                    j += 1
                blocks.append({"type": "table", "headers": headers, "rows": rows})
                i = j
                continue

        # Headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            blocks.append({"type": "heading", "level": len(m.group(1)), "text": _clean_inline(m.group(2))})
            i += 1
            continue

        # Blockquote (assumptions)
        if stripped.startswith(">"):
            blocks.append({"type": "quote", "text": _clean_inline(stripped.lstrip(">").strip())})
            i += 1
            continue

        # Horizontal rule
        if re.fullmatch(r"(-{3,}|_{3,}|\*{3,})", stripped):
            blocks.append({"type": "hr"})
            i += 1
            continue

        # List
        if re.match(r"^(\-|\*|\d+\.)\s+", stripped):
            items = []
            while i < n and re.match(r"^\s*(\-|\*|\d+\.)\s+", lines[i]):
                items.append(_clean_inline(re.sub(r"^\s*(\-|\*|\d+\.)\s+", "", lines[i])))
                i += 1
            blocks.append({"type": "list", "items": items})
            continue

        # Paragraph (join consecutive non-empty, non-special lines)
        para = [stripped]
        i += 1
        while i < n and lines[i].strip() and not re.match(r"^(#{1,6}\s|>|\-\s|\*\s|\d+\.\s|\|)", lines[i].strip()):
            para.append(lines[i].strip())
            i += 1
        blocks.append({"type": "paragraph", "text": _clean_inline(" ".join(para))})
    return blocks


def _kop_lines(meta: Optional[Dict]) -> List[Dict]:
    """Return center/identity rows for kop surat; returns [] when meta not present."""
    if not meta or not meta.get("show_kop"):
        return []
    return [meta]


# ---------------------------------------------------------------------------
# WORD
# ---------------------------------------------------------------------------
def to_docx(title: str, md: str, meta: Optional[Dict]) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    if meta and meta.get("show_kop"):
        for key, size, bold in [("dinas", 12, True), ("nama_sekolah", 14, True)]:
            val = meta.get(key)
            if val:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(val)
                r.bold = bold
                r.font.size = Pt(size)
        addr = meta.get("alamat_sekolah") or ""
        npsn = meta.get("npsn") or ""
        line = addr + (f" | NPSN: {npsn}" if npsn else "")
        if line.strip():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(line).font.size = Pt(10)
        doc.add_paragraph("_" * 70)
        info = [
            ("Mata Pelajaran", "PJOK / Penjasorkes", "Kelas / Fase", meta.get("kelas_fase", "-")),
            ("Materi", meta.get("materi", "-"), "Semester", meta.get("semester", "-")),
            ("Nama Guru", meta.get("nama_guru", "-"), "Tahun Ajaran", meta.get("tahun_ajaran", "-")),
        ]
        t = doc.add_table(rows=0, cols=4)
        for a, b, c, d in info:
            cells = t.add_row().cells
            cells[0].text = a
            cells[1].text = str(b)
            cells[2].text = c
            cells[3].text = str(d)
        doc.add_paragraph("=" * 70)

    doc.add_heading(title, level=0)

    for blk in parse_markdown(md):
        t = blk["type"]
        if t == "heading":
            doc.add_heading(blk["text"], level=min(blk["level"], 4))
        elif t == "paragraph":
            doc.add_paragraph(blk["text"])
        elif t == "quote":
            p = doc.add_paragraph(blk["text"])
            p.runs[0].italic = True if p.runs else None
        elif t == "list":
            for it in blk["items"]:
                doc.add_paragraph(it, style="List Bullet")
        elif t == "hr":
            doc.add_paragraph("_" * 60)
        elif t == "table":
            headers = blk["headers"]
            tbl = doc.add_table(rows=1, cols=len(headers))
            tbl.style = "Table Grid"
            hdr = tbl.rows[0].cells
            for idx, h in enumerate(headers):
                hdr[idx].text = h
                for par in hdr[idx].paragraphs:
                    for run in par.runs:
                        run.bold = True
            for row in blk["rows"]:
                cells = tbl.add_row().cells
                for idx, val in enumerate(row):
                    if idx < len(cells):
                        cells[idx].text = val
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def to_pdf(title: str, md: str, meta: Optional[Dict]) -> bytes:
    bio = io.BytesIO()
    docpdf = SimpleDocTemplate(bio, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                               topMargin=16 * mm, bottomMargin=16 * mm, title=title)
    ss = getSampleStyleSheet()
    normal = ParagraphStyle("body", parent=ss["Normal"], fontSize=10, leading=14)
    center_b = ParagraphStyle("cb", parent=normal, alignment=1, fontSize=13, leading=16)
    center_small = ParagraphStyle("cs", parent=normal, alignment=1, fontSize=9)
    story = []

    if meta and meta.get("show_kop"):
        if meta.get("dinas"):
            story.append(Paragraph(f"<b>{meta['dinas']}</b>", ParagraphStyle("d", parent=center_b, fontSize=11)))
        if meta.get("nama_sekolah"):
            story.append(Paragraph(f"<b>{meta['nama_sekolah']}</b>", center_b))
        addr = (meta.get("alamat_sekolah") or "")
        npsn = meta.get("npsn") or ""
        line = addr + (f" | NPSN: {npsn}" if npsn else "")
        if line.strip():
            story.append(Paragraph(line, center_small))
        story.append(Spacer(1, 4))
        story.append(HRFlowable(width="100%", thickness=1.2, color=colors.black))
        info = [
            ["Mata Pelajaran : PJOK / Penjasorkes", f"Kelas / Fase : {meta.get('kelas_fase','-')}"],
            [f"Materi : {meta.get('materi','-')}", f"Semester : {meta.get('semester','-')}"],
            [f"Nama Guru : {meta.get('nama_guru','-')}", f"Tahun Ajaran : {meta.get('tahun_ajaran','-')}"],
        ]
        it = Table(info, colWidths=["*", "*"])
        it.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 1),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ]))
        story.append(it)
        story.append(HRFlowable(width="100%", thickness=1.2, color=colors.black))
        story.append(Spacer(1, 6))

    story.append(Paragraph(f"<b>{title}</b>", ParagraphStyle("t", parent=normal, fontSize=15, alignment=1)))
    story.append(Spacer(1, 8))

    for blk in parse_markdown(md):
        t = blk["type"]
        if t == "heading":
            fs = max(16 - blk["level"] * 2, 10)
            story.append(Spacer(1, 4))
            story.append(Paragraph(f"<b>{blk['text']}</b>", ParagraphStyle("h", parent=normal, fontSize=fs, leading=fs + 3)))
            story.append(Spacer(1, 2))
        elif t == "paragraph":
            story.append(Paragraph(blk["text"], normal))
            story.append(Spacer(1, 3))
        elif t == "quote":
            story.append(Paragraph(f"<i>{blk['text']}</i>", ParagraphStyle("q", parent=normal, textColor=colors.HexColor("#555555"))))
            story.append(Spacer(1, 3))
        elif t == "list":
            for it in blk["items"]:
                story.append(Paragraph(f"&bull;&nbsp;{it}", normal))
            story.append(Spacer(1, 3))
        elif t == "hr":
            story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#999999")))
            story.append(Spacer(1, 3))
        elif t == "table":
            headers = blk["headers"]
            data = [[Paragraph(f"<b>{h}</b>", ParagraphStyle("th", parent=normal, fontSize=9)) for h in headers]]
            for row in blk["rows"]:
                data.append([Paragraph(c, ParagraphStyle("td", parent=normal, fontSize=9)) for c in row])
            tbl = Table(data, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 1, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAEAEA")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(tbl)
            story.append(Spacer(1, 6))

    docpdf.build(story)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# EXCEL
# ---------------------------------------------------------------------------
def to_xlsx(title: str, md: str, meta: Optional[Dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dokumen"
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = XLFont(bold=True)
    wrap = XLAlign(wrap_text=True, vertical="top")
    hdr_fill = PatternFill("solid", fgColor="EAEAEA")
    r = 1

    def put(text, font=None, merge=6):
        nonlocal r
        c = ws.cell(row=r, column=1, value=text)
        if font:
            c.font = font
        c.alignment = wrap
        if merge > 1:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=merge)
        r += 1

    if meta and meta.get("show_kop"):
        put(meta.get("dinas", ""), XLFont(bold=True, size=12))
        put(meta.get("nama_sekolah", ""), XLFont(bold=True, size=14))
        addr = (meta.get("alamat_sekolah") or "")
        npsn = meta.get("npsn") or ""
        put(addr + (f" | NPSN: {npsn}" if npsn else ""))
        put(f"Kelas / Fase: {meta.get('kelas_fase','-')}  |  Semester: {meta.get('semester','-')}")
        put(f"Guru: {meta.get('nama_guru','-')}  |  Tahun Ajaran: {meta.get('tahun_ajaran','-')}")
        r += 1

    put(title, XLFont(bold=True, size=13))
    r += 1

    for blk in parse_markdown(md):
        t = blk["type"]
        if t == "heading":
            put(blk["text"], bold)
        elif t in ("paragraph", "quote"):
            put(blk["text"])
        elif t == "list":
            for it in blk["items"]:
                put("• " + it)
        elif t == "hr":
            r += 1
        elif t == "table":
            headers = blk["headers"]
            for col, h in enumerate(headers, start=1):
                c = ws.cell(row=r, column=col, value=h)
                c.font = bold
                c.border = border
                c.fill = hdr_fill
                c.alignment = wrap
            r += 1
            for row in blk["rows"]:
                for col, val in enumerate(row, start=1):
                    c = ws.cell(row=r, column=col, value=val)
                    c.border = border
                    c.alignment = wrap
                r += 1
            r += 1
    for col in range(1, 8):
        ws.column_dimensions[chr(64 + col)].width = 22
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
